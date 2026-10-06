"""Delivery pipeline: score -> report -> verify on hand-built records (no area centroids exist yet
for the delivery areas, so OSM cannot be queried for them until derive-centroids runs)."""
import json
from pathlib import Path

import openpyxl
import pytest

from lagosdata.classify import classify
from lagosdata.config import load_config
from lagosdata.pipeline import run_stages
from lagosdata.record import Business
from lagosdata.run import Run
from lagosdata.score import BASIS_FLOOR, BASIS_FULL, COMPONENTS

from .conftest import requires_soffice

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def delivery_run(tmp_path):
    cfg, areas, path = load_config(ROOT / 'config' / 'delivery.yaml')
    run = Run.create(cfg, areas, path, run_id='test-delivery', out_dir=str(tmp_path), echo=False)
    run.cfg.sources.site_contacts.enabled = False      # no network in tests
    run.cfg.enrich.place_pages.enabled = False
    recs = [
        Business(name='Fixture Mart Lekki', label='Supermarket', area='Lekki Phase 1', reviews=900, rating=4.4,
                 website='https://mart.example', phone='+2348031234567', score_basis=BASIS_FULL,
                 additional_info={'Service options': [{'Delivery': True}, {'Takeout': True}]},
                 source='Google Maps (sweep, Oct 2026)'),
        Business(name='Fixture Mart Lekki', label='Supermarket', area='Victoria Island', reviews=300,
                 phone='+2348031234568', source='Google Maps (sweep, Oct 2026)'),
        Business(name='Fixture Pharmacy', label='Pharmacy', area='Yaba', reviews=3,
                 source='OpenStreetMap (Oct 2026)'),
        Business(name='Fixture Buka', label='Restaurant', area='Outside catchment', instagram='fixturebuka',
                 delivery_text_signals=['catering'], source='Google Maps (sweep, Oct 2026)'),
        Business(name='Fixture Chambers', label='Lawyer', area='Yaba', source='Finelib (Oct 2026)'),   # out of scope
    ]
    for r in recs:
        r.sources_all = [r.source]
        r.group, r.sub = classify(r.label, r.name)
    run.write_stage('enrich', recs)
    yield run
    run.close()


def test_scores_and_basis(delivery_run):
    run_stages(delivery_run, ['score'])
    recs = {(r.name, r.area): r for r in delivery_run.read_master()}
    top = recs[('Fixture Mart Lekki', 'Lekki Phase 1')]
    assert top.score_components['Google Delivery flag'] == 25
    assert top.score_components['Chain / multi-branch'] == 10        # same name twice
    assert top.score_basis == BASIS_FULL and top.band.startswith('Hot')
    assert recs[('Fixture Pharmacy', 'Yaba')].score_basis == BASIS_FLOOR   # never presented as a measurement
    assert recs[('Fixture Buka', 'Outside catchment')].score_components['Catering offered'] == 15
    assert ('Fixture Chambers', 'Yaba') not in recs                  # not a food / grocery / pharmacy prospect


def test_workbook_has_every_component_column(delivery_run):
    run_stages(delivery_run, ['score', 'report'])
    wb = openpyxl.load_workbook(delivery_run.workbook_path)
    assert wb.sheetnames == ['DASHBOARD', 'README', 'TARGET LIST', 'SCORING METHOD']
    hdr = [c.value for c in wb['TARGET LIST'][1]]
    assert all(c in hdr for c in COMPONENTS) and 'Score basis' in hdr and 'WhatsApp' in hdr
    assert wb['TARGET LIST']['B2'].value == 'Fixture Mart Lekki'      # sorted by score desc


@requires_soffice
def test_verify_passes(delivery_run):
    assert run_stages(delivery_run, ['score', 'report', 'verify'])
    rep = json.loads((delivery_run.dir / 'verify_report.json').read_text())
    assert rep['ok'] and rep['summary']['score_basis'] == {BASIS_FULL: 1, BASIS_FLOOR: 3}


def test_combine_merges_runs_and_verifies(delivery_run, tmp_path):
    from lagosdata.cli import main
    from lagosdata.combine import combine
    run_stages(delivery_run, ['score'])
    cfg, areas, path = load_config(ROOT / 'config' / 'combined.yaml')
    out = delivery_run.dir.parent
    run = Run.create(cfg, areas, path, run_id='test-combined', out_dir=str(out), echo=False)
    res = combine(run, [delivery_run])
    recs = run.read_master()
    assert res['kept'] == len(recs) == 3                         # the two 'Fixture Mart Lekki' share a name
    mart = next(r for r in recs if r.name == 'Fixture Mart Lekki')
    assert mart.score != '' and mart.priority                    # delivery score kept, magazine priority added
    assert run_stages(run, ['report'])
    import openpyxl
    hdr = [c.value for c in openpyxl.load_workbook(run.workbook_path)['MASTER DATABASE'][1]]
    assert {'Programme', 'Delivery Score', 'Delivery Band'} <= set(hdr)
    run.close()
