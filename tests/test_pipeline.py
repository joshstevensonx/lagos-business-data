"""SPEC §10 step 2-3: prove discover -> geo -> classify -> dedupe -> report -> verify end to end on OSM."""
import json

import openpyxl

from lagosdata.pipeline import run_stages
from lagosdata.stages import verify as V

from .conftest import requires_soffice


def test_end_to_end_without_verify(magazine_run):
    assert run_stages(magazine_run, ['discover', 'geo', 'classify', 'dedupe', 'enrich', 'score', 'report'])
    master = json.loads(magazine_run.master_path.read_text())
    names = [r['name'] for r in master]
    # the supermarket appears as node + way (with 'Ltd'): one record after dedupe
    assert sum('Test Fixture Supermarket Anthony' in n for n in names) == 1
    far = [r for r in master if r['name'] == 'Fixture Supermarket Far Away'][0]
    assert far['area'] == 'Outside catchment' and far['zone'] == 'Outside catchment'
    assert all(r['priority'] for r in master)
    wb = openpyxl.load_workbook(magazine_run.workbook_path)
    assert wb.sheetnames == ['DASHBOARD', 'README', 'MASTER DATABASE', 'AREA SUMMARY', 'GROUP x AREA',
                             'CATEGORY TREE', 'RESEARCH MATRIX', 'SOURCE LOG']
    assert wb['CATEGORY TREE'].max_row >= 4 + 436             # zero rows included
    readme = ' '.join(str(c.value) for row in wb['README'].iter_rows() for c in row if c.value)
    assert '© OpenStreetMap contributors' in readme


def test_rerun_is_a_noop_for_completed_searches(magazine_run, overpass_fixture):
    calls = []
    magazine_run.options['source_kwargs'] = {'osm': {'transport': lambda q: calls.append(q) or overpass_fixture}}
    run_stages(magazine_run, ['discover'])
    run_stages(magazine_run, ['discover'])
    assert len(calls) == 1
    assert len(magazine_run.read_raw('osm')) == len(overpass_fixture['elements'])


@requires_soffice
def test_verify_passes_on_a_good_run(magazine_run):
    assert run_stages(magazine_run, None)
    rep = json.loads((magazine_run.dir / 'verify_report.json').read_text())
    assert rep['ok'], [c for c in rep['checks'] if not c['ok']]
    assert (magazine_run.dir / 'verify_sample.csv').exists()


@requires_soffice
def test_verify_catches_a_kpi_pointing_at_the_wrong_column(magazine_run):
    """The DASHBOARD bug that reached a delivered workbook: 'With a phone number'
    counted the wrong column and read 0. verify must fail on it."""
    run_stages(magazine_run, ['discover', 'geo', 'classify', 'dedupe', 'enrich', 'score', 'report'])
    meta = json.loads((magazine_run.dir / 'workbook_meta.json').read_text())
    kpi = next(k for k in meta['kpis'] if k['key'] == 'phone')
    wb = openpyxl.load_workbook(magazine_run.workbook_path)
    wb['DASHBOARD'][kpi['cell']] = wb['DASHBOARD'][kpi['cell']].value.replace('$P$', '$L$')  # Has Website
    wb.save(magazine_run.workbook_path)
    ok, checks = V.verify(magazine_run)
    assert not ok
    bad = next(c for c in checks if c.name == 'dashboard KPIs recounted in Python')
    assert not bad.ok and 'With a phone number' in bad.detail and 'python=' in bad.detail


def test_verify_without_libreoffice_does_not_claim_success(magazine_run):
    run_stages(magazine_run, ['discover', 'geo', 'classify', 'dedupe', 'enrich', 'score', 'report'])

    def no_soffice(p):
        raise RuntimeError('LibreOffice (soffice) not found')
    ok, checks = V.verify(magazine_run, recalc_fn=no_soffice)
    assert not ok


def test_schema_check_catches_bad_phone(magazine_run):
    run_stages(magazine_run, ['discover', 'geo', 'classify', 'dedupe', 'enrich', 'score', 'report'])
    data = json.loads(magazine_run.master_path.read_text())
    data[0]['phone'] = '+234803123456'                          # 9-digit mobile
    magazine_run.master_path.write_text(json.dumps(data))
    ok, checks = V.verify(magazine_run, recalc_fn=lambda p: p)
    schema = next(c for c in checks if c.name == 'schema')
    assert not schema.ok and 'malformed phones' in schema.detail
