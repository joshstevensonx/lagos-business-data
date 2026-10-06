# -*- coding: utf-8 -*-
"""Delivery prospects workbook - 4 tabs (SPEC §1.2).

Ported from reference/build_delivery_workbook.py. Changes:
  * every score component gets its own column on TARGET LIST (SPEC §1.2)
  * Score basis is 'Full' / 'Floor - not enriched' from the record, never inferred
    from the source string
  * WhatsApp / Twitter-X / TikTok columns for the site-contacts enricher
  * formulas address columns by header name; areas come from config
"""
from __future__ import annotations

import collections

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..record import OUTSIDE, UNASSIGNED, source_kind
from ..score import BASIS_FULL, COMPONENTS
from .style import (FONT, GREY, Cols, blank, body, bold, box, center, fill, kpi_font, sect_fill, sect_font,
                    style_header, style_total_row, sub_font, title_font, widths, write_readme)

TL = "'TARGET LIST'"
CATS = ['Restaurant / Food', 'Supermarket / Grocery', 'Pharmacy', 'Other']
BANDS = ['Hot - approach first', 'Warm - worth a call',
         'Cool - lower priority', 'Cold - little evidence of delivery']
MAXPTS = {'Google Delivery flag': 25, 'No-contact delivery': 5, 'Takeout': 5,
          'Delivery hours published': 5, 'Catering offered': 15,
          'Order-volume proxy (reviews)': 25, 'Has website': 8,
          'Online ordering / menu': 7, 'Chain / multi-branch': 10,
          'Category weight': 10, 'Rating 4.0+': 5}
WHY = {
    'Google Delivery flag': 'The business itself has told Google it delivers. The single strongest direct signal.',
    'No-contact delivery': 'Indicates an established delivery process rather than ad-hoc drop-offs.',
    'Takeout': 'Food already leaves the premises; delivery is a short step from there.',
    'Delivery hours published': 'Separate delivery opening hours imply a real, staffed delivery operation.',
    'Catering offered': 'The best available proxy for genuinely large orders - the volume worth outsourcing.',
    'Order-volume proxy (reviews)': 'Review count stands in for customer throughput. Logarithmic, so a huge chain cannot dominate.',
    'Has website': 'Suggests a business large enough to have invested in its own channel.',
    'Online ordering / menu': 'An order pipeline already exists, which is what a delivery partner plugs into.',
    'Chain / multi-branch': 'Multiple outlets mean multi-drop routes and a single decision-maker for several sites.',
    'Category weight': 'Supermarkets carry the heaviest baskets, restaurants the most frequent drops, pharmacies the most urgent.',
    'Rating 4.0+': 'A well-run business is a better partner and more likely to protect its delivery experience.',
}
COLS = (['Rank', 'Business Name', 'Area', 'Category', 'Google Category', 'Delivery Score', 'Band', 'Score basis']
        + COMPONENTS
        + ['Google Rating', 'Reviews', 'Phone', 'All Phones', 'WhatsApp', 'Website', 'Email', 'Instagram',
           'Facebook', 'Twitter / X', 'LinkedIn', 'TikTok', 'Contact Channels', 'Full Address', 'Latitude',
           'Longitude', 'Google Maps', 'Data Source', 'Last Checked', 'Approach Status', 'Notes'])
APPROACH = 'Not approached,Contacted,Meeting Booked,Proposal Sent,Trial,Won,Declined,Do Not Contact'


def build(recs, cfg, searches, info) -> tuple[Workbook, dict]:
    recs = sorted(recs, key=lambda r: (-float(r.score or 0), r.area, r.name.lower()))
    N = len(recs)
    LAST = max(N + 1, 2)
    TODAY = info['today']
    AREAS = list(cfg.areas) + [OUTSIDE, UNASSIGNED]
    C = Cols(COLS)
    meta = {'pipeline': 'delivery', 'rows': N, 'kpis': [], 'area_table': [], 'checks': {}}
    wb = Workbook()

    def R(h, absolute=True):
        return C.rng(TL, h, LAST, absolute)

    # ---------------------------------------------------------------- TARGET LIST
    ws = wb.create_sheet('TARGET LIST')
    for i, h in enumerate(COLS, start=1): ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(COLS))
    band_fills = {BANDS[0]: ('C6EFCE', '006100'), BANDS[1]: ('FFF2CC', '7F6000'),
                  BANDS[2]: ('FCE4D6', '833C0B'), BANDS[3]: (GREY, '595959')}
    for i, r in enumerate(recs):
        rr = i + 2
        comps = r.score_components or {}
        vals = ([i + 1, r.name, r.area, r.delivery_category, r.label, r.score, r.band, r.score_basis]
                + [comps.get(k, 0) for k in COMPONENTS]
                + [r.rating if r.rating != '' else None, r.reviews if r.reviews != '' else None,
                   r.phone, r.phones_all, r.whatsapp, r.website, r.email, r.instagram, r.facebook, r.twitter,
                   r.linkedin, r.tiktok, r.contact_channels, r.addr, r.lat if r.lat != '' else None,
                   r.lng if r.lng != '' else None, r.maps, r.source, TODAY, 'Not approached', r.notes])
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=rr, column=c, value=blank(v))
            cell.font = body
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (C.n('Full Address'), C.n('Notes'))))
        ws.cell(row=rr, column=C.n('Google Rating')).number_format = '0.0'
        ws.cell(row=rr, column=C.n('Reviews')).number_format = '#,##0'
        ws.cell(row=rr, column=C.n('Latitude')).number_format = '0.000000'
        ws.cell(row=rr, column=C.n('Longitude')).number_format = '0.000000'
        ws.cell(row=rr, column=C.n('Delivery Score')).number_format = '0'
        f, col = band_fills.get(r.band, band_fills[BANDS[3]])
        for h in ('Delivery Score', 'Band'):
            cell = ws.cell(row=rr, column=C.n(h))
            cell.fill = fill(f)
            cell.font = Font(name=FONT, size=10, color=col, bold=(r.band == BANDS[0]))
        if r.score_basis != BASIS_FULL:
            ws.cell(row=rr, column=C.n('Score basis')).font = Font(name=FONT, size=10, italic=True, color='833C0B')
    widths(ws, [6, 38, 22, 22, 24, 13, 28, 20] + [11] * len(COMPONENTS)
           + [10, 9, 17, 26, 17, 32, 30, 32, 32, 24, 32, 24, 11, 44, 12, 12, 30, 26, 12, 18, 30])
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{LAST}'
    dv = DataValidation(type='list', allow_blank=True, formula1=f'"{APPROACH}"')
    ws.add_data_validation(dv); dv.add(f'{C["Approach Status"]}2:{C["Approach Status"]}{LAST}')

    # ----------------------------------------------------------------- DASHBOARD
    ws = wb.create_sheet('DASHBOARD', 0)
    ws.sheet_view.showGridLines = False
    ws['A1'] = 'LAGOS DELIVERY PROSPECTS'; ws['A1'].font = title_font
    ws['A2'] = (f'Restaurants, supermarkets and pharmacies across {len(cfg.areas)} Lagos areas, ranked by how '
                f'likely they are to need a dedicated delivery partner  |  Built {TODAY}')
    ws['A2'].font = sub_font
    ws.merge_cells('A1:F1'); ws.merge_cells('A2:F2')
    ws['A4'] = 'HEADLINE NUMBERS'; ws['A4'].font = sect_font; ws['A4'].fill = sect_fill
    ws.merge_cells('A4:F4')
    kpis = [
        ('total', 'Total prospects', f'=COUNTA({R("Business Name", False)})'),
        ('hot', BANDS[0], f'=COUNTIF({R("Band")},"{BANDS[0]}")'),
        ('warm', BANDS[1], f'=COUNTIF({R("Band")},"{BANDS[1]}")'),
        ('cool', BANDS[2], f'=COUNTIF({R("Band")},"{BANDS[2]}")'),
        ('cold', BANDS[3], f'=COUNTIF({R("Band")},"{BANDS[3]}")'),
        ('delivery_flag', 'Flagged by Google as doing delivery', f'=COUNTIF({R("Google Delivery flag")},">0")'),
        ('catering', 'Offer catering (bulk-order signal)', f'=COUNTIF({R("Catering offered")},">0")'),
        ('chain', 'Chains / multi-branch', f'=COUNTIF({R("Chain / multi-branch")},">0")'),
        ('phone', 'With a phone number', f'=COUNTIF({R("Phone")},"<>")'),
        ('whatsapp', 'With WhatsApp', f'=COUNTIF({R("WhatsApp")},"<>")'),
        ('website', 'With a website', f'=COUNTIF({R("Website")},"<>")'),
        ('email', 'With an email', f'=COUNTIF({R("Email")},"<>")'),
        ('instagram', 'With Instagram', f'=COUNTIF({R("Instagram")},"<>")'),
        ('full_basis', 'Delivery data actually checked (Score basis = Full)', f'=COUNTIF({R("Score basis")},"{BASIS_FULL}")'),
    ]
    r0 = 5
    for i, (key, lab, f) in enumerate(kpis):
        ws.cell(row=r0 + i, column=1, value=lab).font = bold
        c = ws.cell(row=r0 + i, column=3, value=f)
        c.font = kpi_font; c.alignment = Alignment(horizontal='right'); c.number_format = '#,##0'
        ws.cell(row=r0 + i, column=1).border = box; c.border = box
        meta['kpis'].append({'key': key, 'label': lab, 'sheet': 'DASHBOARD', 'cell': f'C{r0 + i}', 'formula': f})

    ar = r0 + len(kpis) + 1
    ws.cell(row=ar, column=1, value='BY AREA AND CATEGORY').font = sect_font
    ws.cell(row=ar, column=1).fill = sect_fill
    ws.merge_cells(start_row=ar, start_column=1, end_row=ar, end_column=7)
    hdr = ['Area'] + CATS + ['Total', 'Hot']
    for i, h in enumerate(hdr, start=1): ws.cell(row=ar + 1, column=i, value=h)
    style_header(ws, ar + 1, len(hdr))
    first = ar + 2
    A, CAT, B = R('Area'), R('Category'), R('Band')
    for i, a in enumerate(AREAS):
        rr = first + i
        ws.cell(row=rr, column=1, value=a).font = bold
        for j, cat in enumerate(CATS):
            ws.cell(row=rr, column=2 + j, value=f'=COUNTIFS({A},$A{rr},{CAT},"{cat}")')
        ws.cell(row=rr, column=6, value=f'=SUM(B{rr}:E{rr})')
        ws.cell(row=rr, column=7, value=f'=COUNTIFS({A},$A{rr},{B},"{BANDS[0]}")')
        for c in range(1, 8):
            cell = ws.cell(row=rr, column=c); cell.border = box
            if c > 1: cell.font = body; cell.alignment = center
            if i % 2: cell.fill = fill(GREY)
        meta['area_table'].append({'area': a, 'row': rr})
    tot = first + len(AREAS)
    ws.cell(row=tot, column=1, value='TOTAL')
    for c in range(2, 8):
        L = get_column_letter(c)
        ws.cell(row=tot, column=c, value=f'=SUM({L}{first}:{L}{tot - 1})')
    style_total_row(ws, tot, 7)
    meta['area_table_cols'] = {**{cat: get_column_letter(2 + j) for j, cat in enumerate(CATS)},
                               'total': 'F', 'hot': 'G'}
    meta['checks']['area_table_total'] = {'sheet': 'DASHBOARD', 'cell': f'F{tot}'}
    widths(ws, [48, 20, 22, 14, 10, 10, 10])
    ws.freeze_panes = 'A5'

    # ------------------------------------------------------------ SCORING METHOD
    ws = wb.create_sheet('SCORING METHOD')
    ws.sheet_view.showGridLines = False
    ws['A1'] = 'HOW THE DELIVERY SCORE IS BUILT'; ws['A1'].font = title_font
    ws['A2'] = ('Google has no "does large deliveries" field. This score combines the signals that do exist, so '
                'the ranking is a judgement aid, not a measurement. Every component is shown so you can disagree '
                'with it.')
    ws['A2'].font = sub_font
    ws.merge_cells('A1:D1'); ws.merge_cells('A2:D2'); ws.row_dimensions[2].height = 30
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    hdr = ['Signal', 'Max points', 'Why it matters', 'How many prospects have it']
    for i, h in enumerate(hdr, start=1): ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, 4)
    counts = collections.Counter()
    for r in recs:
        for k, v in (r.score_components or {}).items():
            if v: counts[k] += 1
    for i, comp in enumerate(COMPONENTS, start=5):
        ws.cell(row=i, column=1, value=comp).font = bold
        ws.cell(row=i, column=2, value=MAXPTS.get(comp, 0)).alignment = center
        c = ws.cell(row=i, column=3, value=WHY.get(comp, '')); c.font = body
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(row=i, column=4, value=counts.get(comp, 0)).alignment = center
        for cc in range(1, 5): ws.cell(row=i, column=cc).border = box
        ws.row_dimensions[i].height = 34
    br = 5 + len(COMPONENTS) + 1
    ws.cell(row=br, column=1, value='BANDS').font = sect_font
    ws.cell(row=br, column=1).fill = sect_fill
    ws.merge_cells(start_row=br, start_column=1, end_row=br, end_column=4)
    for i, (b, rng) in enumerate(zip(BANDS, ['60 and above', '40 - 59', '25 - 39', 'under 25']), start=br + 1):
        ws.cell(row=i, column=1, value=b).font = bold
        ws.cell(row=i, column=2, value=rng).alignment = center
        ws.cell(row=i, column=4, value=sum(1 for r in recs if r.band == b)).alignment = center
        for cc in range(1, 5): ws.cell(row=i, column=cc).border = box
    widths(ws, [32, 14, 72, 24])

    # ------------------------------------------------------------------- README
    ws = wb.create_sheet('README')
    write_readme(ws, 'ABOUT THIS LIST', readme_rows(recs, cfg, searches, info))
    widths(ws, [22, 30, 24, 24, 24])

    del wb['Sheet']
    wb._sheets = [wb[n] for n in ['DASHBOARD', 'README', 'TARGET LIST', 'SCORING METHOD']]
    return wb, meta


def readme_rows(recs, cfg, searches, info):
    n_full = sum(1 for r in recs if r.score_basis == BASIS_FULL)
    n_floor = len(recs) - n_full
    srcs = sorted({s for r in recs for s in (r.sources_all or [r.source]) if s})
    done = sum(1 for s in searches if s['status'] == 'done')
    not_built = info.get('sources_not_built') or []
    rows = [
        ('Purpose', 'A prospecting list for approaching restaurants, supermarkets and pharmacies that move enough '
                    'volume to need a dedicated delivery partner.'),
        ('Areas covered', ', '.join(cfg.areas) + f'. Prospects found beyond these are kept as "{OUTSIDE}".'),
        ('Categories', 'Restaurants and food outlets, supermarkets and grocery stores, and pharmacies, searched '
                       f'with {len(cfg.taxonomy.terms or [])} terms so the long tail is included.'),
        ('How to use it', 'The TARGET LIST is sorted by Delivery Score, highest first. Start at the top of the Hot '
                          'band, work down, and record outcomes in Approach Status. Filter by Area to plan a day of '
                          'visits in one part of Lagos.'),
        ('Scoring', 'See the SCORING METHOD tab. Every component that fed each score is its own column on the '
                    'TARGET LIST, so you can re-sort on whatever you trust most.'),
        ('What the score is NOT', f'{n_full} prospects show "Full" under Score basis: their Google place page was '
                                  f'checked for delivery, takeout and catering. The other {n_floor} show "Floor - '
                                  'not enriched": they were scored on reviews, website, chain status, category and '
                                  'rating alone and could not earn the 55 points tied to delivery, catering and '
                                  'takeout signals. A low score on a Floor row is a floor, not a verdict - it may '
                                  'simply be unchecked.'),
        ('Contact details', 'Phone comes from the listing. Website, email, WhatsApp and social handles come from the '
                            'business\'s own website where one exists; many Nigerian SMEs trade through Instagram '
                            'and WhatsApp rather than email.'),
        ('Data source', f'{done} searches. ' + (' / '.join(srcs) if srcs else 'none')
                        + (f'. Configured but not yet built, so not run: {", ".join(not_built)}.' if not_built else '')),
        ('Honest limitations', 'Delivery capability is inferred, not confirmed. Google\'s delivery flag is '
                               'self-reported and patchy in Nigeria, so a "No" is weak evidence of absence. Review '
                               'counts proxy for footfall, not delivery volume. Treat the score as a call order and '
                               'confirm volumes on the phone.'),
        ('Conduct and data protection', 'Free sources only, read at low volume; no CAPTCHA solved, no login '
                                        'bypassed, robots.txt honoured. Business contact details for B2B use under '
                                        'Nigeria\'s NDPA - honour every opt-out (Approach Status "Do Not Contact").'),
    ]
    if any(source_kind(s) == 'osm' for s in srcs):
        rows.append(('Attribution', 'Rows sourced from OpenStreetMap: © OpenStreetMap contributors, ODbL.'))
    return rows
