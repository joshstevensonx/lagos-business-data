"""Shared workbook styling - Arial, NAVY 1F3864, BLUE 2F5597 (as delivered)."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
NAVY = '1F3864'; BLUE = '2F5597'; LIGHT = 'D9E2F3'; GREY = 'F2F2F2'; GOLD = 'FFF2CC'
hdr_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
hdr_fill = PatternFill('solid', fgColor=NAVY)
title_font = Font(name=FONT, bold=True, size=16, color=NAVY)
sub_font = Font(name=FONT, italic=True, size=10, color='595959')
sect_font = Font(name=FONT, bold=True, size=11, color='FFFFFF')
sect_fill = PatternFill('solid', fgColor=BLUE)
body = Font(name=FONT, size=10)
bold = Font(name=FONT, size=10, bold=True)
kpi_font = Font(name=FONT, size=11, bold=True, color=NAVY)
total_font = Font(name=FONT, bold=True, color='FFFFFF')
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center')


def fill(color):
    return PatternFill('solid', fgColor=color)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = box


def widths(ws, w):
    for i, x in enumerate(w, start=1):
        ws.column_dimensions[get_column_letter(i)].width = x


def style_total_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c); cell.fill = fill(BLUE)
        cell.border = box; cell.font = total_font
        cell.alignment = center
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='left')


def blank(v):
    """'' and None both become a truly empty cell, so COUNTIF(...,"<>") counts honestly."""
    return None if v == '' else v


class Cols:
    """Header-name -> column letter. Formulas use names, never hand-typed letters:
    a column move then cannot silently re-point a KPI (the 'Has Website' vs
    'Has Phone' bug)."""

    def __init__(self, headers):
        self.headers = list(headers)
        if len(set(self.headers)) != len(self.headers):
            raise ValueError('duplicate column headers')
        self.idx = {h: i + 1 for i, h in enumerate(self.headers)}

    def __getitem__(self, h) -> str:
        return get_column_letter(self.idx[h])

    def n(self, h) -> int:
        return self.idx[h]

    def rng(self, sheet: str, h: str, last: int, absolute: bool = True) -> str:
        L = self[h]
        return f'{sheet}!${L}$2:${L}${last}' if absolute else f'{sheet}!{L}2:{L}{last}'


def write_readme(ws, title, rows, width_cols=5):
    ws.sheet_view.showGridLines = False
    ws['A1'] = title; ws['A1'].font = title_font
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols + 1)
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = bold
        ws.cell(row=i, column=1).fill = fill(LIGHT)
        ws.cell(row=i, column=1).alignment = Alignment(vertical='top', wrap_text=True)
        c = ws.cell(row=i, column=2, value=v); c.font = body
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=width_cols + 1)
        ws.row_dimensions[i].height = max(30, 15 * (len(v) // 110 + 1))
        ws.cell(row=i, column=1).border = box; c.border = box
