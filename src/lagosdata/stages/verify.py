"""verify: independent recount of the delivered workbook (SPEC §9). Non-negotiable.

Every dashboard KPI is recomputed in Python straight from master.json and compared
with the value LibreOffice computes from the workbook's formula. A formula
pointing at the wrong column (the 'Has Website' vs 'Has Phone' bug that read 0
against a true 631) fails here, loudly, with both numbers and the formula text.
"""
from __future__ import annotations

import collections
import csv
import dataclasses
import json
import random
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

from ..record import (CORE_ZONE, OUTSIDE, RING_ZONE, UNASSIGNED, UNCLASSIFIED, Business, source_kind)
from ..score import BASIS_FULL
from ..tree import ALL_SUBS, GROUPS

ERROR_VALUES = ('#REF!', '#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#NUM!', '#NULL!', 'Err:')
PHONE_RE = re.compile(r'^\+234\d{10}$|^\+2341\d{7,8}$')
MULTI = 'Multiple sources (cross-verified)'


@dataclasses.dataclass
class Check:
    name: str
    ok: bool
    detail: str = ''


# ----------------------------------------------------------------- recalc ---
def recalc(xlsx: Path) -> Path:
    """Recalculate every formula with LibreOffice headless; return the recalculated copy."""
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        raise RuntimeError('LibreOffice (soffice) not found - verify cannot compute formula values. '
                           'Install LibreOffice; verification is not claimed without it.')
    tmp = Path(tempfile.mkdtemp(prefix='lagosdata-recalc-'))
    profile = tmp / 'profile'
    out = tmp / 'out'
    cmd = [soffice, f'-env:UserInstallation=file://{profile}', '--headless', '--calc',
           '--convert-to', 'xlsx', '--outdir', str(out), str(xlsx)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=600)
    res = out / xlsx.name
    if not res.exists():
        raise RuntimeError(f'LibreOffice produced no output for {xlsx}')
    return res


def scan_errors(wb) -> list[str]:
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERROR_VALUES):
                    bad.append(f'{ws.title}!{c.coordinate}={c.value}')
    return bad


# -------------------------------------------------------- python recounts ---
def magazine_expected(recs: list[Business]) -> dict:
    z = collections.Counter(r.zone for r in recs)
    pr = collections.Counter(r.priority for r in recs)
    return {
        'total': len(recs),
        'core': z[CORE_ZONE], 'ring': z[RING_ZONE], 'outside': z[OUTSIDE], 'unassigned': z[UNASSIGNED],
        'phone': sum(1 for r in recs if r.contactable == 'Yes'),
        'google_verified': sum(1 for r in recs if r.verification == 'Google-verified'),
        'prio_a': pr['A - High commercial relevance'],
        'prio_b': pr['B - Potential advertiser'],
        'prio_c': pr['C - Ring area, contactable'],
        'cross_verified': sum(1 for r in recs if r.source == MULTI),
    }


def magazine_group_expected(recs, g) -> dict:
    inside = [r for r in recs if r.group == g and r.zone in (CORE_ZONE, RING_ZONE)]
    core = sum(1 for r in inside if r.zone == CORE_ZONE)
    ring = len(inside) - core
    return {'core': core, 'ring': ring, 'total': core + ring,
            'phone': sum(1 for r in inside if r.contactable == 'Yes'),
            'prio_a': sum(1 for r in recs if r.group == g and r.priority == 'A - High commercial relevance')}


def delivery_expected(recs: list[Business]) -> dict:
    b = collections.Counter(r.band for r in recs)
    comp = lambda k: sum(1 for r in recs if (r.score_components or {}).get(k, 0) > 0)  # noqa: E731
    return {
        'total': len(recs),
        'hot': b['Hot - approach first'], 'warm': b['Warm - worth a call'],
        'cool': b['Cool - lower priority'], 'cold': b['Cold - little evidence of delivery'],
        'delivery_flag': comp('Google Delivery flag'), 'catering': comp('Catering offered'),
        'chain': comp('Chain / multi-branch'),
        'phone': sum(1 for r in recs if r.phone), 'whatsapp': sum(1 for r in recs if r.whatsapp),
        'website': sum(1 for r in recs if r.website), 'email': sum(1 for r in recs if r.email),
        'instagram': sum(1 for r in recs if r.instagram),
        'full_basis': sum(1 for r in recs if r.score_basis == BASIS_FULL),
    }


def delivery_area_expected(recs, area, cats) -> dict:
    rs = [r for r in recs if r.area == area]
    out = {c: sum(1 for r in rs if r.delivery_category == c) for c in cats}
    out['total'] = sum(out[c] for c in cats)
    out['hot'] = sum(1 for r in rs if r.band == 'Hot - approach first')
    return out


def _val(wb, sheet, cell):
    v = wb[sheet][cell].value
    return 0 if v is None else v


# ------------------------------------------------------------------ verify ---
def verify(run, recalc_fn=recalc) -> tuple[bool, list[Check]]:
    checks: list[Check] = []
    cfg = run.cfg
    raw = json.loads(run.master_path.read_text())
    recs = [Business.from_dict(d) for d in raw]
    meta = json.loads((run.dir / 'workbook_meta.json').read_text())
    N = len(recs)

    if N == 0:
        checks.append(Check('records present', False, 'master.json is empty - nothing to deliver'))

    # 2. recalc + error scan -------------------------------------------------
    try:
        rec_path = recalc_fn(run.workbook_path)
        wb = openpyxl.load_workbook(rec_path, data_only=True)
        errs = scan_errors(wb)
        checks.append(Check('formula errors (LibreOffice recalc)', not errs,
                            f'errors_found={len(errs)}' + (': ' + ', '.join(errs[:10]) if errs else '')))
    except Exception as e:  # noqa: BLE001
        checks.append(Check('formula errors (LibreOffice recalc)', False, f'could not recalculate: {e}'))
        wb = None

    # 1. KPI recount ---------------------------------------------------------
    if wb is not None:
        exp = magazine_expected(recs) if meta['pipeline'] == 'magazine' else delivery_expected(recs)
        bad = []
        for k in meta['kpis']:
            if k['key'] not in exp:
                bad.append(f'{k["label"]}: no independent Python recount exists for this KPI')
                continue
            got = _val(wb, k['sheet'], k['cell'])
            if got != exp[k['key']]:
                bad.append(f'{k["label"]} ({k["cell"]}): workbook={got} python={exp[k["key"]]} '
                           f'formula={k["formula"]}')
        formulas = openpyxl.load_workbook(run.workbook_path)['DASHBOARD']
        if meta['pipeline'] == 'magazine':
            cols = meta['group_table_cols']
            for gt in meta['group_table']:
                e = magazine_group_expected(recs, gt['group'])
                for key, L in cols.items():
                    got = _val(wb, 'DASHBOARD', f'{L}{gt["row"]}')
                    if got != e[key]:
                        ref = f'{L}{gt["row"]}'
                        bad.append(f'group "{gt["group"]}" {key} ({ref}): workbook={got} python={e[key]} '
                                   f'formula={formulas[ref].value}')
        else:
            cols = meta['area_table_cols']
            cats = [c for c in cols if c not in ('total', 'hot')]
            for at in meta['area_table']:
                e = delivery_area_expected(recs, at['area'], cats)
                for key, L in cols.items():
                    got = _val(wb, 'DASHBOARD', f'{L}{at["row"]}')
                    if got != e[key]:
                        ref = f'{L}{at["row"]}'
                        bad.append(f'area "{at["area"]}" {key} ({ref}): workbook={got} python={e[key]} '
                                   f'formula={formulas[ref].value}')
        n_checked = len(meta['kpis']) + len(meta.get('group_table') or meta.get('area_table') or [])
        checks.append(Check('dashboard KPIs recounted in Python', not bad,
                            f'{n_checked} KPI rows checked' + ('; MISMATCH: ' + ' | '.join(bad) if bad else '')))

    # 3. schema --------------------------------------------------------------
    all_fields = [f.name for f in dataclasses.fields(Business)]
    nones = [(i, f) for i, d in enumerate(raw) for f in all_fields if f in d and d[f] is None]
    bad_phone = [r.phone for r in recs if r.phone and not PHONE_RE.match(r.phone)]
    groups = set(GROUPS)
    subs = set(ALL_SUBS) | {UNCLASSIFIED}
    bad_group = [r.group for r in recs if r.group not in groups]
    bad_sub = [(r.group, r.sub) for r in recs if (r.group, r.sub) not in subs]
    areas_ok = set(cfg.areas) | {OUTSIDE, UNASSIGNED}
    bad_area = [r.area for r in recs if r.area not in areas_ok]
    problems = []
    if nones: problems.append(f'{len(nones)} None values (first: record {nones[0][0]} .{nones[0][1]})')
    if bad_phone: problems.append(f'{len(bad_phone)} malformed phones e.g. {bad_phone[:3]}')
    if bad_group: problems.append(f'{len(bad_group)} unknown groups e.g. {sorted(set(bad_group))[:3]}')
    if bad_sub: problems.append(f'{len(bad_sub)} (group, sub) not in tree e.g. {sorted(set(bad_sub))[:3]}')
    if bad_area: problems.append(f'{len(bad_area)} unknown areas e.g. {sorted(set(bad_area))[:3]}')
    checks.append(Check('schema', not problems, '; '.join(problems) or f'{N} records OK'))

    # 4. unclassified share --------------------------------------------------
    unc = [r for r in recs if (r.group, r.sub) == UNCLASSIFIED]
    share = len(unc) / N if N else 0
    top = collections.Counter(r.label or '(no label)' for r in unc).most_common(20)
    checks.append(Check('unclassified < 5%', share < 0.05,
                        f'{len(unc)}/{N} = {share:.1%}' + (f'; top labels: {top}' if share >= 0.05 else '')))

    # 5. dedupe sanity -------------------------------------------------------
    by_phone = collections.defaultdict(list)
    for r in recs:
        if r.phone:
            by_phone[r.phone].append(r)
    dup_p, shared_ok = [], []
    for p, rs in by_phone.items():
        if len(rs) < 2:
            continue
        ids = [r.place_id for r in rs]
        # Google itself lists these as distinct places (e.g. two branches on one hotline): the
        # place-id key outranks phone (SPEC §6.3), so this is not a dedupe failure - but say so
        if all(ids) and len(set(ids)) == len(ids):
            shared_ok.append(p)
        else:
            dup_p.append(p)
    dup_id = [p for p, n in collections.Counter(r.place_id for r in recs if r.place_id).items() if n > 1]
    detail = (f'shared phones: {dup_p[:5]}; shared place_ids: {dup_id[:5]}' if dup_p or dup_id
              else 'no duplicate phone or place_id')
    if shared_ok:
        detail += (f'; {len(shared_ok)} phone(s) shared by distinct Google places (branches/one owner), '
                   f'kept separate: {shared_ok[:5]}')
    checks.append(Check('dedupe sanity', not dup_p and not dup_id, detail))

    # 6. row counts across tabs ----------------------------------------------
    if wb is not None:
        sheet = 'MASTER DATABASE' if meta['pipeline'] == 'magazine' else 'TARGET LIST'
        rows_in_sheet = sum(1 for row in wb[sheet].iter_rows(min_row=2, max_col=2, values_only=True) if row[1])
        counts = {'master.json': N, sheet: rows_in_sheet}
        for name, loc in meta['checks'].items():
            if isinstance(loc, dict):
                counts[name] = _val(wb, loc['sheet'], loc['cell'])
        checks.append(Check('row counts agree across tabs', len(set(counts.values())) == 1, json.dumps(counts)))

    # 7. sample for human eyeballing ------------------------------------------
    rnd = random.Random(run.id)
    sample = rnd.sample(recs, min(10, N))
    path = run.dir / 'verify_sample.csv'
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh)
        w.writerow(['name', 'area', 'group', 'sub', 'phone', 'addr', 'source', 'maps', 'checked_by',
                    'matches_listing (Y/N)', 'notes'])
        for r in sample:
            w.writerow([r.name, r.area, r.group, r.sub, r.phone, r.addr, r.source, r.maps, '', '', ''])
    checks.append(Check('spot-check sample written', True,
                        f'{len(sample)} records in {path.name} - for a human to check against the listing; '
                        'NOT verified by this tool'))

    ok = all(c.ok for c in checks)
    (run.dir / 'verify_report.json').write_text(json.dumps(
        {'ok': ok, 'checks': [dataclasses.asdict(c) for c in checks], 'summary': summary(recs, cfg)},
        indent=1, ensure_ascii=False))
    return ok, checks


# 8. summary ------------------------------------------------------------------
def summary(recs: list[Business], cfg) -> dict:
    by_source = collections.Counter()
    for r in recs:
        for s in set(r.sources_all or [r.source]):
            by_source[source_kind(s)] += 1
    out = {
        'records': len(recs),
        'by_source_kind': dict(by_source.most_common()),
        'by_area': dict(collections.Counter(r.area for r in recs).most_common()),
        'by_group': dict(collections.Counter(r.group for r in recs).most_common()),
        'contactable': sum(1 for r in recs if r.phone),
        'enrichment': {k: sum(1 for r in recs if getattr(r, k))
                       for k in ('website', 'email', 'whatsapp', 'instagram', 'facebook')},
    }
    if cfg.pipeline == 'delivery':
        out['score_basis'] = dict(collections.Counter(r.score_basis for r in recs))
        out['bands'] = dict(collections.Counter(r.band for r in recs))
    return out


def format_summary(s: dict) -> str:
    lines = [f'records: {s["records"]}   contactable (phone): {s["contactable"]}']
    for k in ('by_source_kind', 'by_area', 'by_group', 'enrichment', 'score_basis', 'bands'):
        if k in s:
            lines.append(f'{k}:')
            lines += [f'  {n:>6}  {name}' for name, n in s[k].items()]
    return '\n'.join(lines)
